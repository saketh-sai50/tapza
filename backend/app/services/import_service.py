import openpyxl
from io import BytesIO
from typing import Dict, Any, List, Tuple
from app.db import transaction

REQUIRED_HEADERS = ["name", "mrp"]
EXPECTED_COLUMNS = [
    "name",
    "generic_name",
    "manufacturer",
    "category",
    "dosage_form",
    "pack_size",
    "mrp",
    "requires_prescription"
]

def parse_and_validate_excel(file_stream: BytesIO, db) -> Dict[str, Any]:
    """
    Parses an uploaded .xlsx workbook, rigorously validates headers and rows,
    and performs atomic batch insertion/upsert into the medicines table.
    """
    try:
        workbook = openpyxl.load_workbook(file_stream, data_only=True)
    except Exception as e:
        raise ValueError(f"Invalid Excel file format: {str(e)}")

    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        raise ValueError("Excel file is empty or missing data rows (first row must be header)")

    raw_headers = rows[0]
    normalized_headers = [str(h).strip().lower().replace(" ", "_") if h is not None else "" for h in raw_headers]

    # Validate essential headers exist
    for req in REQUIRED_HEADERS:
        if req not in normalized_headers:
            raise ValueError(f"Missing required column header: '{req.upper()}'")

    col_index = {name: idx for idx, name in enumerate(normalized_headers) if name}

    imported_count = 0
    updated_count = 0
    errors: List[Dict[str, Any]] = []
    valid_items: List[Dict[str, Any]] = []

    for row_num, row_data in enumerate(rows[1:], start=2):
        if not any(row_data):
            # Skip completely empty rows
            continue

        def get_val(col_name: str, default=None):
            idx = col_index.get(col_name)
            if idx is not None and idx < len(row_data):
                val = row_data[idx]
                return val if val is not None else default
            return default

        # 1. Validate Medicine Name
        raw_name = get_val("name")
        name = str(raw_name).strip() if raw_name else ""
        if not name:
            errors.append({"row": row_num, "error": "Medicine Name cannot be empty"})
            continue

        # 2. Validate MRP (Maximum Retail Price)
        raw_mrp = get_val("mrp")
        try:
            mrp = float(raw_mrp)
            if mrp <= 0:
                raise ValueError("MRP must be greater than zero")
        except (ValueError, TypeError):
            errors.append({"row": row_num, "medicine": name, "error": f"Invalid MRP value: '{raw_mrp}'. Must be a positive number."})
            continue

        # 3. Optional metadata
        generic_name = str(get_val("generic_name", "")).strip() or None
        manufacturer = str(get_val("manufacturer", "")).strip() or None
        category = str(get_val("category", "General")).strip() or "General"
        dosage_form = str(get_val("dosage_form", "")).strip() or None
        pack_size = str(get_val("pack_size", "1 Unit")).strip() or "1 Unit"

        # 4. Prescription requirement flag
        raw_rx = get_val("requires_prescription", 0)
        requires_rx = 0
        if isinstance(raw_rx, bool):
            requires_rx = 1 if raw_rx else 0
        elif isinstance(raw_rx, (int, float)):
            requires_rx = 1 if int(raw_rx) == 1 else 0
        elif isinstance(raw_rx, str):
            requires_rx = 1 if raw_rx.strip().lower() in ("yes", "y", "true", "1") else 0

        valid_items.append({
            "name": name,
            "generic_name": generic_name,
            "manufacturer": manufacturer,
            "category": category,
            "dosage_form": dosage_form,
            "pack_size": pack_size,
            "mrp": round(mrp, 2),
            "requires_prescription": requires_rx,
            "row_num": row_num
        })

    # Atomic persistence in database
    with transaction(db):
        for item in valid_items:
            existing = db.execute(
                "SELECT id FROM medicines WHERE LOWER(name) = LOWER(?)",
                (item["name"],)
            ).fetchone()

            if existing:
                db.execute(
                    """
                    UPDATE medicines
                    SET generic_name = ?, manufacturer = ?, category = ?, dosage_form = ?,
                        pack_size = ?, mrp = ?, requires_prescription = ?, is_active = 1
                    WHERE id = ?
                    """,
                    (
                        item["generic_name"],
                        item["manufacturer"],
                        item["category"],
                        item["dosage_form"],
                        item["pack_size"],
                        item["mrp"],
                        item["requires_prescription"],
                        existing["id"]
                    )
                )
                updated_count += 1
            else:
                db.execute(
                    """
                    INSERT INTO medicines (name, generic_name, manufacturer, category, dosage_form, pack_size, mrp, requires_prescription, is_active)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, 1)
                    """,
                    (
                        item["name"],
                        item["generic_name"],
                        item["manufacturer"],
                        item["category"],
                        item["dosage_form"],
                        item["pack_size"],
                        item["mrp"],
                        item["requires_prescription"]
                    )
                )
                imported_count += 1

    return {
        "total_rows_processed": len(rows) - 1,
        "imported": imported_count,
        "updated": updated_count,
        "rejected_count": len(errors),
        "errors": errors
    }


def generate_catalogue_template() -> BytesIO:
    """Generates an Excel (.xlsx) starter template with headers and demo rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Medicine Catalogue"

    headers = [
        "Name",
        "Generic Name",
        "Manufacturer",
        "Category",
        "Dosage Form",
        "Pack Size",
        "MRP",
        "Requires Prescription"
    ]
    ws.append(headers)

    sample_rows = [
        ["Paracetamol 500mg", "Acetaminophen", "GlaxoSmithKline", "Analgesic", "Tablet", "10 tablets", 15.50, "No"],
        ["Amoxicillin 500mg", "Amoxicillin Trihydrate", "Novartis", "Antibiotic", "Capsule", "10 capsules", 85.00, "Yes"],
        ["Cetirizine 10mg", "Cetirizine Hydrochloride", "Dr. Reddy's", "Antihistamine", "Tablet", "10 tablets", 22.00, "No"],
        ["Metformin 500mg", "Metformin Hydrochloride", "Sun Pharma", "Antidiabetic", "Tablet", "20 tablets", 45.00, "Yes"],
        ["Omeprazole 20mg", "Omeprazole", "Cipla", "Antacid", "Capsule", "14 capsules", 55.00, "No"]
    ]
    for row in sample_rows:
        ws.append(row)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
