from flask import Blueprint, request, jsonify, g
from app.db import get_db, transaction
from app.utils.auth_helpers import token_required, role_required

admin_bp = Blueprint("admin", __name__, url_prefix="/api/admin")

@admin_bp.before_request
@token_required
@role_required("ADMIN")
def enforce_admin_only():
    """Guarantees all endpoints in this blueprint strictly require ADMIN role."""
    pass


@admin_bp.route("/metrics", methods=["GET"])
def get_platform_metrics():
    """
    Returns platform-wide KPIs:
      - Total orders, active broadcasts, confirmed, delivered, cancelled
      - Registered users by role
      - Active vs inactive pharmacies
      - Platform GMV (Gross Merchandise Value)
      - Zone distribution
    """
    db = get_db()

    # User metrics
    total_users = db.execute("SELECT COUNT(*) as c FROM users").fetchone()["c"]
    total_customers = db.execute("SELECT COUNT(*) as c FROM users WHERE role = 'CUSTOMER'").fetchone()["c"]
    total_pharmacies = db.execute("SELECT COUNT(*) as c FROM users WHERE role = 'PHARMACY'").fetchone()["c"]
    active_pharmacies = db.execute("SELECT COUNT(*) as c FROM pharmacy_profiles WHERE is_active = 1").fetchone()["c"]

    # Order metrics
    total_orders = db.execute("SELECT COUNT(*) as c FROM orders").fetchone()["c"]
    active_broadcasts = db.execute("SELECT COUNT(*) as c FROM orders WHERE status = 'BROADCASTING'").fetchone()["c"]
    confirmed_orders = db.execute("SELECT COUNT(*) as c FROM orders WHERE status IN ('CONFIRMED', 'PREPARING', 'DISPATCHED', 'DELIVERED')").fetchone()["c"]
    delivered_orders = db.execute("SELECT COUNT(*) as c FROM orders WHERE status = 'DELIVERED'").fetchone()["c"]
    expired_orders = db.execute("SELECT COUNT(*) as c FROM orders WHERE status = 'EXPIRED'").fetchone()["c"]
    cancelled_orders = db.execute("SELECT COUNT(*) as c FROM orders WHERE status = 'CANCELLED'").fetchone()["c"]

    # Quote metrics
    total_quotes = db.execute("SELECT COUNT(*) as c FROM quotes").fetchone()["c"]
    full_quotes = db.execute("SELECT COUNT(*) as c FROM quotes WHERE is_full_fulfillment = 1").fetchone()["c"]
    partial_quotes = db.execute("SELECT COUNT(*) as c FROM quotes WHERE is_full_fulfillment = 0").fetchone()["c"]

    # Platform GMV (sum of final_amount on confirmed/delivered orders)
    gmv_row = db.execute(
        """
        SELECT COALESCE(SUM(final_amount), 0) as gmv
        FROM orders
        WHERE status IN ('CONFIRMED', 'PREPARING', 'DISPATCHED', 'DELIVERED')
        """
    ).fetchone()
    platform_gmv = round(float(gmv_row["gmv"]), 2)

    # Zone distribution
    zone_rows = db.execute(
        """
        SELECT zone,
               COUNT(*) as order_count,
               COALESCE(SUM(final_amount), 0) as zone_volume
        FROM orders
        GROUP BY zone
        ORDER BY order_count DESC
        """
    ).fetchall()
    zone_stats = [
        {
            "zone": zr["zone"],
            "order_count": zr["order_count"],
            "zone_volume": round(float(zr["zone_volume"]), 2)
        }
        for zr in zone_rows
    ]

    return jsonify({
        "metrics": {
            "users": {
                "total": total_users,
                "customers": total_customers,
                "pharmacies": total_pharmacies,
                "active_pharmacies": active_pharmacies
            },
            "orders": {
                "total": total_orders,
                "active_broadcasts": active_broadcasts,
                "confirmed": confirmed_orders,
                "delivered": delivered_orders,
                "expired": expired_orders,
                "cancelled": cancelled_orders
            },
            "quotes": {
                "total": total_quotes,
                "full_fulfillment": full_quotes,
                "partial_fulfillment": partial_quotes
            },
            "platform_gmv": platform_gmv,
            "zones": zone_stats
        }
    }), 200


@admin_bp.route("/pharmacies", methods=["GET"])
def list_pharmacies():
    """Lists all onboarded pharmacies with license and performance metadata."""
    db = get_db()
    rows = db.execute(
        """
        SELECT u.id as user_id, u.email, u.full_name, u.phone, u.created_at,
               p.id as profile_id, p.pharmacy_name, p.license_number, p.address, p.zone, p.is_active,
               (SELECT COUNT(*) FROM quotes WHERE pharmacy_id = u.id) as quotes_submitted_count,
               (SELECT COUNT(*) FROM order_assignments WHERE pharmacy_id = u.id) as orders_won_count
        FROM users u
        JOIN pharmacy_profiles p ON u.id = p.user_id
        WHERE u.role = 'PHARMACY'
        ORDER BY p.is_active DESC, u.id DESC
        """
    ).fetchall()

    pharmacies = [
        {
            "user_id": r["user_id"],
            "profile_id": r["profile_id"],
            "email": r["email"],
            "full_name": r["full_name"],
            "phone": r["phone"],
            "pharmacy_name": r["pharmacy_name"],
            "license_number": r["license_number"],
            "address": r["address"],
            "zone": r["zone"],
            "is_active": bool(r["is_active"]),
            "quotes_submitted_count": r["quotes_submitted_count"],
            "orders_won_count": r["orders_won_count"],
            "created_at": str(r["created_at"])
        }
        for r in rows
    ]
    return jsonify({"pharmacies": pharmacies}), 200


@admin_bp.route("/pharmacies/<int:pharmacy_id>/toggle", methods=["POST"])
def toggle_pharmacy_status(pharmacy_id: int):
    """Enables or suspends a pharmacy from receiving order broadcasts."""
    db = get_db()
    pharmacy = db.execute(
        "SELECT id, pharmacy_name, is_active FROM pharmacy_profiles WHERE user_id = ?",
        (pharmacy_id,)
    ).fetchone()

    if not pharmacy:
        return jsonify({"error": "Pharmacy not found"}), 404

    new_status = 0 if pharmacy["is_active"] else 1
    with transaction(db):
        db.execute(
            "UPDATE pharmacy_profiles SET is_active = ? WHERE user_id = ?",
            (new_status, pharmacy_id)
        )

    return jsonify({
        "message": f"Pharmacy '{pharmacy['pharmacy_name']}' is now {'active' if new_status else 'suspended'}",
        "pharmacy_id": pharmacy_id,
        "is_active": bool(new_status)
    }), 200


@admin_bp.route("/pharmacies/<int:pharmacy_id>/zone", methods=["PUT"])
def update_pharmacy_zone(pharmacy_id: int):
    """Reassigns a pharmacy's operating broadcast zone."""
    data = request.get_json() or {}
    new_zone = (data.get("zone") or "").strip().upper()

    if not new_zone:
        return jsonify({"error": "Zone is required (e.g. ZONE_A, ZONE_B, CENTRAL)"}), 400

    db = get_db()
    with transaction(db):
        db.execute("UPDATE pharmacy_profiles SET zone = ? WHERE user_id = ?", (new_zone, pharmacy_id))
        db.execute("UPDATE users SET zone = ? WHERE id = ?", (new_zone, pharmacy_id))

    return jsonify({
        "message": f"Pharmacy zone updated to {new_zone}",
        "pharmacy_id": pharmacy_id,
        "zone": new_zone
    }), 200


@admin_bp.route("/orders", methods=["GET"])
def list_all_orders():
    """
    Global order audit log across all zones and statuses.
    Supports filtering by status and zone.
    """
    db = get_db()
    status_filter = request.args.get("status")
    zone_filter = request.args.get("zone")

    query = """
        SELECT o.id, o.customer_id, o.zone, o.delivery_address, o.status,
               o.broadcast_started_at, o.broadcast_expires_at,
               o.total_mrp, o.final_amount, o.mock_payment_status, o.created_at,
               u.full_name as customer_name, u.email as customer_email,
               oa.pharmacy_id, p.pharmacy_name as assigned_pharmacy_name,
               oa.confirmation_status as pharmacy_confirmation_status,
               (SELECT COUNT(*) FROM quotes WHERE order_id = o.id) as quote_count,
               (SELECT COUNT(*) FROM order_items WHERE order_id = o.id) as item_count
        FROM orders o
        JOIN users u ON o.customer_id = u.id
        LEFT JOIN order_assignments oa ON o.id = oa.order_id
        LEFT JOIN pharmacy_profiles p ON oa.pharmacy_id = p.user_id
        WHERE 1=1
    """
    params = []

    if status_filter:
        query += " AND o.status = ?"
        params.append(status_filter)
    if zone_filter:
        query += " AND o.zone = ?"
        params.append(zone_filter)

    query += " ORDER BY o.id DESC LIMIT 100"

    rows = db.execute(query, params).fetchall()

    orders = [
        {
            "id": r["id"],
            "customer": {
                "id": r["customer_id"],
                "name": r["customer_name"],
                "email": r["customer_email"]
            },
            "zone": r["zone"],
            "delivery_address": r["delivery_address"],
            "status": r["status"],
            "total_mrp": float(r["total_mrp"]),
            "final_amount": float(r["final_amount"]) if r["final_amount"] is not None else float(r["total_mrp"]),
            "mock_payment_status": r["mock_payment_status"],
            "item_count": r["item_count"],
            "quote_count": r["quote_count"],
            "assigned_pharmacy": {
                "id": r["pharmacy_id"],
                "name": r["assigned_pharmacy_name"],
                "confirmation_status": r["pharmacy_confirmation_status"]
            } if r["pharmacy_id"] else None,
            "broadcast_started_at": str(r["broadcast_started_at"]),
            "broadcast_expires_at": str(r["broadcast_expires_at"]),
            "created_at": str(r["created_at"])
        }
        for r in rows
    ]

    return jsonify({"orders": orders}), 200


@admin_bp.route("/pharmacies/import", methods=["POST"])
def import_pharmacies_excel():
    """
    Validated Excel (.xlsx) bulk onboarding of zone-wise pharmacy credentials.
    Reads: Pharmacy Name, Email, Password, Zone, License Number, Address, Phone, Pharmacist Name.
    Validates email format, unique credentials, required zones, and hashes passwords securely.
    """
    from io import BytesIO
    import openpyxl
    from app.utils.auth_helpers import hash_password

    if "file" not in request.files:
        return jsonify({"error": "No file uploaded. Please upload with key 'file'"}), 400

    uploaded = request.files["file"]
    if not uploaded.filename or not uploaded.filename.lower().endswith(".xlsx"):
        return jsonify({"error": "Please upload a valid Excel (.xlsx) file"}), 400

    try:
        wb = openpyxl.load_workbook(BytesIO(uploaded.read()), data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2:
            return jsonify({"error": "Excel sheet is empty or missing headers"}), 400

        headers = [str(h).strip().lower().replace(" ", "_") if h else "" for h in rows[0]]
        col_map = {name: idx for idx, name in enumerate(headers) if name}

        required_cols = ["pharmacy_name", "email", "password", "zone", "license_number"]
        for rc in required_cols:
            if rc not in col_map:
                return jsonify({"error": f"Missing required column: '{rc.replace('_', ' ').title()}'"}), 400

        db = get_db()
        imported = 0
        updated = 0
        errors = []

        with transaction(db):
            for row_idx, r in enumerate(rows[1:], start=2):
                if not any(r):
                    continue

                def val(k, d=""):
                    idx = col_map.get(k)
                    if idx is not None and idx < len(r) and r[idx] is not None:
                        return str(r[idx]).strip()
                    return d

                p_name = val("pharmacy_name")
                email = val("email").lower()
                pwd = val("password")
                zone = val("zone").upper()
                license_no = val("license_number")
                address = val("address", "Sector Main Road")
                phone = val("phone", "+1-800-PHARMA")
                full_name = val("pharmacist_name") or val("full_name") or f"{p_name} Pharmacist"

                if not email or "@" not in email:
                    errors.append({"row": row_idx, "error": f"Invalid email: '{email}'"})
                    continue
                if not pwd or len(pwd) < 6:
                    errors.append({"row": row_idx, "error": "Password must be at least 6 characters"})
                    continue
                if not zone:
                    errors.append({"row": row_idx, "error": "Zone is required"})
                    continue

                existing = db.execute("SELECT id FROM users WHERE email = ?", (email,)).fetchone()
                pwd_h = hash_password(pwd)

                if existing:
                    u_id = existing["id"]
                    db.execute(
                        "UPDATE users SET password_hash = ?, full_name = ?, phone = ?, zone = ? WHERE id = ?",
                        (pwd_h, full_name, phone, zone, u_id)
                    )
                    db.execute(
                        """
                        UPDATE pharmacy_profiles
                        SET pharmacy_name = ?, license_number = ?, address = ?, zone = ?, is_active = 1
                        WHERE user_id = ?
                        """,
                        (p_name, license_no, address, zone, u_id)
                    )
                    updated += 1
                else:
                    cur = db.execute(
                        """
                        INSERT INTO users (email, password_hash, full_name, role, phone, zone)
                        VALUES (?, ?, ?, 'PHARMACY', ?, ?)
                        """,
                        (email, pwd_h, full_name, phone, zone)
                    )
                    u_id = cur.lastrowid
                    db.execute(
                        """
                        INSERT INTO pharmacy_profiles (user_id, pharmacy_name, license_number, address, zone, is_active)
                        VALUES (?, ?, ?, ?, ?, 1)
                        """,
                        (u_id, p_name, license_no, address, zone)
                    )
                    imported += 1

        return jsonify({
            "message": "Pharmacy credentials import processed",
            "stats": {
                "imported": imported,
                "updated": updated,
                "errors": errors
            }
        }), 200
    except Exception as exc:
        return jsonify({"error": f"Import failed: {str(exc)}"}), 500


@admin_bp.route("/pharmacies/template", methods=["GET"])
def download_pharmacy_template():
    """Generates a downloadable starter Excel template for zone-wise pharmacy onboarding."""
    import openpyxl
    from io import BytesIO
    from flask import send_file

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pharmacy Credentials"

    headers = [
        "Pharmacy Name", "Email", "Password", "Zone",
        "License Number", "Address", "Phone", "Pharmacist Name"
    ]
    ws.append(headers)

    samples = [
        ["City Medicos Zone A", "city.zonea@tapza.com", "CityPass@123", "ZONE_A", "LIC-ZA-101", "12 Healthcare Blvd, Zone A", "+1-555-1011", "Dr. Alan Miller"],
        ["Metro Health Zone B", "metro.zoneb@tapza.com", "MetroPass@123", "ZONE_B", "LIC-ZB-202", "45 Metro Center, Zone B", "+1-555-2022", "Dr. Clara Barton"],
        ["Central Core Pharmacy", "central.core@tapza.com", "CentralPass@123", "CENTRAL", "LIC-ZC-303", "88 Downtown Cross, Central", "+1-555-3033", "Dr. David Ross"]
    ]
    for s in samples:
        ws.append(s)

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)

    return send_file(
        bio,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="tapza_pharmacies_template.xlsx"
    )


@admin_bp.route("/users", methods=["GET"])
def list_users():
    """Lists platform users for administrator management."""
    db = get_db()
    rows = db.execute(
        "SELECT id, email, full_name, role, phone, zone, created_at FROM users ORDER BY id DESC"
    ).fetchall()

    users = [
        {
            "id": r["id"],
            "email": r["email"],
            "full_name": r["full_name"],
            "role": r["role"],
            "phone": r["phone"],
            "zone": r["zone"],
            "created_at": str(r["created_at"])
        }
        for r in rows
    ]
    return jsonify({"users": users}), 200
